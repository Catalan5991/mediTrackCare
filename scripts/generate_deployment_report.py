from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "entregable" / "Reporte_despliegue_MediTrackCare.docx"
NAVY = "17365D"
BLUE = "2E74B5"
LIGHT = "E8EEF5"
GRAY = "666666"


def font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    if shd.getparent() is None:
        tc_pr.append(shd)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}")) or OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        if node.getparent() is None:
            tc_mar.append(node)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for old in list(grid):
        grid.remove(old)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, label in enumerate(headers):
        shade(hdr.cells[i], LIGHT)
        p = hdr.cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        font(r, 9.5, True, NAVY)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            font(r, 9.5)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(text)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F4F7")
    pPr.append(shd)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
    r._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
    r.font.size = Pt(8.5)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    font(run, 9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15
for name, size, before, after, color in (
    ("Heading 1", 16, 18, 10, BLUE),
    ("Heading 2", 13, 12, 6, BLUE),
    ("Heading 3", 12, 8, 4, NAVY),
):
    st = styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
for name in ("List Bullet", "List Number"):
    styles[name].font.name = "Calibri"
    styles[name].font.size = Pt(11)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = header.add_run("MediTrackCare | Liberación del servicio web")
font(r, 9, True, GRAY)
add_page_number(section.footer.paragraphs[0])

# Portada editorial
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(90)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("UNIVERSIDAD TECNOLÓGICA DE\nIZÚCAR DE MATAMOROS")
font(r, 13, True, NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(34)
p.paragraph_format.space_after = Pt(10)
r = p.add_run("REPORTE DE DESPLIEGUE\nDEL SERVICIO WEB")
font(r, 25, True, NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MediTrackCare API")
font(r, 18, True, BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
r = p.add_run("Desarrollo Web Integral | Parcial 3 | Evaluación ordinaria")
font(r, 11, color=GRAY)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(72)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(
    "Integrantes:\n"
    "Lisbeth Catalan Castro — 23292062\n"
    "Emmanuel Castillo Arias — 2329XXXX\n"
    "Orlando Sánchez — 2329XXXX\n"
    "Grupo: 9 B\n"
    "Docente: MTE. Alfonso Felipe Lima Cortés\n"
    "Fecha de entrega: 25 de agosto de 2026"
)
font(r, 11)
doc.add_page_break()

doc.add_heading("Introducción", level=1)
doc.add_paragraph(
    "El presente reporte explica desde cero cómo publicar MediTrackCare, un servicio web creado con NestJS, Prisma y PostgreSQL. El trabajo se realiza en dos etapas. Primero se configura Docker Desktop en una computadora con Windows y se levantan la API y la base de datos con Docker Compose. Después se usa el mismo Dockerfile para publicar el backend en Railway. Los pasos se explican con palabras sencillas, pero también se incluyen los comandos, las variables, las pruebas y las evidencias necesarias para comprobar que todo funciona."
)
doc.add_paragraph(
    "La idea principal es que el proyecto pueda instalarse de la misma manera en cualquier equipo. Docker reúne la aplicación y sus dependencias dentro de contenedores. Cuando la API inicia, Prisma aplica los cambios de la base de datos y carga la información inicial. Esa carga puede ejecutarse varias veces sin repetir los mismos datos. Las contraseñas se guardan como variables de entorno y la ruta /api/health permite revisar rápidamente si el servicio está activo."
)

doc.add_heading("1. Descripción del proyecto", level=1)
doc.add_paragraph(
    "MediTrackCare es una API orientada al seguimiento integral de tratamientos médicos. Centraliza usuarios y roles, perfiles de pacientes, médicos y cuidadores, medicamentos, recetas e indicaciones, recordatorios, registro de tomas, mediciones médicas, notificaciones, archivos, bitácora de auditoría y tableros. La API expone sus recursos bajo el prefijo /api y ofrece documentación interactiva Swagger en /docs."
)
add_table(doc, ["Componente", "Tecnología", "Responsabilidad"], [
    ("API", "NestJS 11 / TypeScript", "Controladores, validación, autenticación JWT y reglas de acceso."),
    ("Persistencia", "PostgreSQL 17", "Datos clínicos, usuarios, relaciones y auditoría."),
    ("ORM", "Prisma 7", "Migraciones, cliente tipado y acceso a PostgreSQL."),
    ("Distribución", "Docker Desktop / Compose", "Contenedores Linux ejecutados desde Windows y reutilizados en Railway."),
    ("Documentación", "Swagger", "Exploración y prueba manual de endpoints."),
], [1800, 2160, 5400])

doc.add_heading("2. Despliegue y pruebas", level=1)
doc.add_heading("2.1 Estado inicial y criterios de aceptación", level=2)
doc.add_paragraph(
    "El punto de partida es una computadora con Windows que todavía no tiene Docker ni PostgreSQL ejecutándose. El 24 de agosto de 2026 se probó la ruta de salud iniciando NestJS directamente en localhost y se obtuvo {\"status\":\"ok\",\"service\":\"MediTrackCare API\",\"timestamp\":\"2026-08-25T01:25:20.954Z\"}. Esto demuestra que la ruta funciona, pero todavía falta probarla dentro de Docker. El despliegue local se considerará terminado cuando los contenedores api y db estén activos, y Railway se considerará terminado cuando el dominio público responda con código HTTP 200."
)
add_table(doc, ["Instancia", "Objetivo", "Criterio de aceptación", "Estado inicial"], [
    ("Windows local", "API + PostgreSQL en Docker Desktop", "Dos servicios activos; health 200; Swagger accesible; datos persistentes", "Pendiente de Docker"),
    ("Railway", "API desde Dockerfile + PostgreSQL administrado", "Build correcto; migraciones/semilla; dominio y health 200", "Pendiente"),
], [1900, 2800, 3260, 1400])

doc.add_heading("2.2 Despliegue local desde cero en Windows", level=2)
doc.add_heading("2.2.1 Requisitos de la computadora", level=3)
doc.add_paragraph("Se necesita Windows 10 u 11 de 64 bits, conexión a Internet, virtualización activada y al menos 8 GB de memoria RAM. Docker Desktop utiliza WSL 2 para ejecutar contenedores Linux dentro de Windows. Aunque el trabajo se realiza desde PowerShell, la API y PostgreSQL se ejecutan dentro de contenedores Linux, de forma parecida a un servidor Linux.")
add_code(doc, "winver\nGet-ComputerInfo | Select-Object WindowsProductName, WindowsVersion, OsArchitecture\nwsl --version\nwsl --status")
doc.add_paragraph("Si wsl --version no muestra información, se abre PowerShell como administrador y se instala o actualiza WSL. Después puede ser necesario reiniciar la computadora:")
add_code(doc, "wsl --install\nwsl --update\nwsl --set-default-version 2\nwsl --list --online\n# Instalar Ubuntu 22.04 si aparece con ese nombre en la lista:\nwsl --install -d Ubuntu-22.04")
doc.add_paragraph("Al abrir Ubuntu 22.04 por primera vez, Windows solicita un nombre de usuario y una contraseña para Linux. Después, en Docker Desktop se entra a Settings > Resources > WSL Integration y se activa la integración con Ubuntu 22.04. De esta manera la práctica conserva un entorno Linux aunque los comandos principales se ejecuten desde Windows.")
doc.add_paragraph("En el Administrador de tareas, en la pestaña Rendimiento > CPU, debe aparecer Virtualización: habilitada. Si aparece deshabilitada, se debe activar Intel VT-x o AMD-V desde BIOS/UEFI antes de continuar.")

doc.add_heading("2.2.2 Descarga e instalación de Docker Desktop", level=3)
doc.add_paragraph("Entrar a https://www.docker.com/products/docker-desktop/ o a la guía oficial de instalación. Seleccionar Download for Windows y descargar Docker Desktop Installer.exe. No se recomienda descargar el instalador desde páginas de terceros.")
for text in (
    "Cerrar programas importantes antes de iniciar la instalación.",
    "Hacer doble clic en Docker Desktop Installer.exe.",
    "Elegir la instalación por usuario, que es la opción recomendada para la mayoría de los estudiantes.",
    "Mantener seleccionada la opción para usar WSL 2 como motor de Linux.",
    "Finalizar el asistente y reiniciar Windows si lo solicita.",
    "Buscar Docker Desktop en el menú Inicio, abrirlo y aceptar los términos de uso.",
    "Esperar hasta que la aplicación muestre Docker Desktop is running.",
):
    add_number(doc, text)
doc.add_paragraph("La instalación también puede iniciarse desde PowerShell, ubicándose en la carpeta donde se descargó el archivo:")
add_code(doc, "cd $HOME\\Downloads\nStart-Process '.\\Docker Desktop Installer.exe' -Wait -ArgumentList 'install', '--user'")
doc.add_paragraph("En Docker Desktop se revisa Settings > General y se confirma Use the WSL 2 based engine. En la parte inferior izquierda debe indicar que se están usando Linux containers.")
add_code(doc, "docker version\ndocker compose version\ndocker run --rm hello-world")
doc.add_paragraph("Si docker version solo muestra la parte Client o informa que no puede conectarse, Docker Desktop todavía no terminó de iniciar. Se debe abrir la aplicación, esperar unos minutos y repetir la prueba.")

doc.add_heading("2.2.3 Obtención y revisión del proyecto", level=3)
add_code(doc, "git clone https://github.com/Catalan5991/mediTrackCare.git meditrackcare\ncd meditrackcare\ngit status\ngit log -1 --oneline")
doc.add_paragraph("Antes de construir deben estar presentes Dockerfile, docker-compose.yml, railway.toml, prisma/schema.prisma, prisma/migrations, prisma/seed.ts y .env.production.example. No se debe copiar el archivo .env de desarrollo.")
add_code(doc, "test -f Dockerfile && test -f docker-compose.yml && test -f prisma/schema.prisma\ndocker compose config --no-interpolate")

doc.add_heading("2.2.4 Variables y secretos de producción", level=3)
doc.add_paragraph("Crear .env.production desde el ejemplo, restringir sus permisos y sustituir cada valor. Las contraseñas y cadenas reales no deben aparecer en Git, Swagger, capturas o este reporte.")
add_code(doc, "Copy-Item .env.production.example .env.production\n# Abrir .env.production con Visual Studio Code o Bloc de notas.\ncode .env.production")
doc.add_paragraph("Para crear un JWT_SECRET se puede usar Node.js. Se copia el resultado dentro de .env.production y no se comparte en capturas:")
add_code(doc, "node -e \"console.log(require('crypto').randomBytes(48).toString('base64'))\"")
add_table(doc, ["Variable", "Uso", "Regla"], [
    ("POSTGRES_DB", "Nombre de la base local", "meditrackcare o valor equivalente"),
    ("POSTGRES_USER", "Usuario PostgreSQL", "No usar postgres en producción"),
    ("POSTGRES_PASSWORD", "Contraseña de PostgreSQL", "Larga, única y no versionada"),
    ("DATABASE_URL", "Conexión Prisma", "Compose la reemplaza para usar el host db"),
    ("JWT_SECRET", "Firma de tokens", "Aleatorio, mínimo 32 caracteres"),
    ("ADMIN_EMAIL / ADMIN_PASSWORD", "Administrador inicial", "Correo válido; contraseña mínima de 12 caracteres"),
    ("SEED_DEMO_DATA", "Carga académica", "true solo sin datos reales; false en producción"),
    ("API_PORT", "Puerto del host", "3000 para localhost"),
], [2300, 3700, 3360])

doc.add_heading("2.2.5 Construcción, migración y arranque", level=3)
doc.add_paragraph("El servicio db crea PostgreSQL 17 y conserva /var/lib/postgresql/data en un volumen. api se construye en dos etapas, espera la salud de la base y ejecuta prisma migrate deploy, prisma db seed y node dist/main mediante npm run start:deploy.")
add_code(doc, "docker compose pull\ndocker compose build --no-cache api\ndocker compose up -d\ndocker compose ps\ndocker compose logs --tail=200 api")
doc.add_paragraph("En los registros deben aparecer las migraciones aplicadas, la semilla completada y el arranque de NestJS sin excepciones. La semilla garantiza ADMIN, DOCTOR, CAREGIVER y PATIENT; el registro público asigna PATIENT y solo ADMIN administra usuarios y roles.")

doc.add_heading("2.2.6 Verificación local", level=3)
add_code(doc, "curl.exe -i http://localhost:3000/api/health\ncurl.exe -I http://localhost:3000/docs\ndocker inspect --format='{{json .State.Health}}' meditrackcare-api-1\ndocker volume ls | Select-String meditrackcare")
doc.add_paragraph("La URL de Swagger es http://localhost:3000/docs. En Windows puede aparecer una solicitud del firewall; solo se debe permitir acceso en redes privadas si otros equipos necesitan hacer pruebas. PostgreSQL no se publica hacia Windows, por lo que permanece dentro de la red de Docker.")

doc.add_heading("2.3 Pruebas a la instancia Docker local", level=2)
doc.add_paragraph("Primero se revisa que los contenedores estén activos y después se prueban las funciones de la API. Los códigos de respuesta también pueden verse desde Swagger. Para entrar a rutas protegidas se inicia sesión como administrador y se envía Authorization: Bearer TOKEN. El token debe ocultarse en las capturas.")
add_code(doc, "curl.exe -s http://localhost:3000/api/health\ncurl.exe -X POST http://localhost:3000/api/auth/login -H 'Content-Type: application/json' -d '{\"email\":\"CORREO_ADMIN\",\"password\":\"CONTRASENA\"}'\ncurl.exe -i http://localhost:3000/api/auth/profile -H 'Authorization: Bearer TOKEN'")
add_table(doc, ["ID", "Prueba", "Esperado", "Resultado/evidencia"], [
    ("LOC-01", "GET /api/health", "HTTP 200 y status ok", "Ejecutar en contenedor"),
    ("LOC-02", "Swagger /docs", "Interfaz accesible", "Ejecutar en contenedor"),
    ("LOC-03", "Login ADMIN", "Token y rol ADMIN", "Pendiente"),
    ("LOC-04", "Ruta protegida sin token", "HTTP 401", "Pendiente"),
    ("LOC-05", "PATIENT intenta administrar roles", "HTTP 403", "Pendiente"),
    ("LOC-06", "CRUD clínico", "Códigos conforme a Swagger", "Pendiente"),
    ("LOC-07", "Reinicio de Compose", "Datos permanecen", "Pendiente"),
    ("LOC-08", "Pruebas unitarias", "40/40 aprobadas", "Verificado 24/08/2026"),
    ("LOC-09", "E2E en servidor", "Todas aprobadas", "Aún no ejecutadas"),
], [1100, 3150, 2600, 2510])
add_code(doc, "npm ci\nnpx prisma generate\nnpm run test:db:migrate\nnpm run test:e2e")
doc.add_paragraph("Las E2E usan TEST_DATABASE_URL y nunca deben apuntar a la base de producción. El registro histórico del proyecto reporta 157/157 pruebas E2E aprobadas el 17/07/2026; debe anexarse una ejecución nueva de la instancia preparada para esta entrega.")

doc.add_heading("2.4 Persistencia, respaldo y actualización local", level=2)
add_code(doc, "docker compose exec -T db pg_dump -U meditrack -d meditrackcare -Fc > meditrackcare_2026-08-24.dump\nGet-Item .\\meditrackcare_2026-08-24.dump")
doc.add_paragraph("El nombre de usuario y base debe coincidir con .env.production. El respaldo contiene datos sensibles y debe almacenarse cifrado. Para actualizar: obtener cambios, construir una nueva imagen, levantar Compose y revisar salud. No usar docker compose down -v salvo que se quiera eliminar deliberadamente la base.")
add_code(doc, "git pull\ndocker compose build api\ndocker compose up -d\ndocker compose ps\ndocker compose logs --tail=100 api")

doc.add_heading("2.5 Despliegue desde cero en Railway", level=2)
doc.add_paragraph("Railway no ejecuta Compose como una sola unidad: el PostgreSQL administrado y el backend son servicios distintos. El repositorio debe estar publicado en GitHub. Railway detecta el Dockerfile ubicado en la raíz; el archivo railway.toml declara /api/health y la política de reinicio.")
for text in (
    "Crear una cuenta y un proyecto vacío en Railway.",
    "Agregar un servicio PostgreSQL desde + New > Database > PostgreSQL.",
    "Agregar un servicio desde el repositorio GitHub y seleccionar la rama de entrega.",
    "En Variables del backend crear DATABASE_URL como referencia a ${{Postgres.DATABASE_URL}}; adaptar Postgres al nombre real del servicio.",
    "Agregar NODE_ENV=production, JWT_SECRET, JWT_EXPIRES_IN_SECONDS, ADMIN_NAME, ADMIN_EMAIL, ADMIN_PASSWORD y SEED_DEMO_DATA=false.",
    "No fijar PORT: Railway lo inyecta y NestJS ya escucha process.env.PORT.",
    "En Settings comprobar Dockerfile, configurar Healthcheck Path=/api/health y desplegar.",
    "Revisar Deploy Logs hasta observar migraciones, semilla y arranque; después generar un dominio en Networking.",
):
    add_number(doc, text)
add_table(doc, ["Variable Railway", "Valor recomendado"], [
    ("DATABASE_URL", "${{Postgres.DATABASE_URL}} mediante referencia, no copiar credenciales"),
    ("NODE_ENV", "production"),
    ("JWT_SECRET", "Secreto aleatorio diferente al local"),
    ("JWT_EXPIRES_IN_SECONDS", "3600"),
    ("ADMIN_NAME", "Administrador MediTrackCare"),
    ("ADMIN_EMAIL", "Correo administrativo del equipo"),
    ("ADMIN_PASSWORD", "Contraseña única de al menos 12 caracteres"),
    ("SEED_DEMO_DATA", "false; true solo en entorno académico aislado"),
], [3000, 6360])
doc.add_paragraph("La base administrada permanece privada. Para este backend no es necesario habilitar acceso público a PostgreSQL. La URL pública del backend y la fecha efectiva se incorporarán después del primer despliegue satisfactorio.")

doc.add_heading("2.6 Pruebas a la instancia Railway", level=2)
add_code(doc, "curl.exe -i https://DOMINIO_RAILWAY/api/health\ncurl.exe -I https://DOMINIO_RAILWAY/docs")
add_table(doc, ["ID", "Solicitud", "Resultado esperado", "Resultado real"], [
    ("VPS-01", "GET /api/health", "200; status=ok", "Pendiente de URL"),
    ("VPS-02", "GET /docs", "Swagger accesible", "Pendiente"),
    ("VPS-03", "POST /api/auth/login", "Token ADMIN", "Pendiente"),
    ("VPS-04", "GET /api/auth/profile con token", "200 y perfil", "Pendiente"),
    ("VPS-05", "Ruta protegida sin token", "401", "Pendiente"),
    ("VPS-06", "Operación sin permiso", "403", "Pendiente"),
    ("VPS-07", "CRUD de módulos", "Conforme a Swagger", "Pendiente"),
], [1200, 3000, 2900, 2260])

doc.add_heading("2.7 Diagnóstico de fallas frecuentes", level=2)
add_table(doc, ["Síntoma", "Causa probable", "Acción"], [
    ("API reinicia", "DATABASE_URL, migración o secreto ausente", "Revisar logs y variables; no imprimir secretos"),
    ("P1001/Pg inaccesible", "Base no saludable o URL incorrecta", "Esperar db local; verificar referencia PostgreSQL en Railway"),
    ("Healthcheck falla", "Aplicación no escucha PORT o arranque excede tiempo", "Revisar PORT, /api/health y logs de migración"),
    ("401", "Token ausente, expirado o usuario/rol inactivo", "Iniciar sesión y comprobar estado del usuario y rol"),
    ("403", "Rol sin permiso requerido", "Consultar matriz; no elevar permisos para ocultar un error"),
    ("Rol no asignable", "Rol fuera del catálogo o desactivado", "Usar ADMIN, DOCTOR, CAREGIVER o PATIENT activo"),
    ("Datos desaparecen", "Volumen no montado o se usó down -v", "Verificar volumen y restaurar respaldo"),
], [2300, 3300, 3760])

doc.add_heading("3. Alcances logrados", level=1)
for text in (
    "Backend modular de NestJS con prefijo global, validación de DTO, filtros y autenticación JWT.",
    "Persistencia PostgreSQL definida mediante Prisma y migraciones versionadas.",
    "Imagen de producción y orquestación local con comprobaciones de salud y volumen persistente.",
    "Carga inicial que puede ejecutarse varias veces sin duplicar roles ni usuarios de ejemplo.",
    "Catálogo de roles homologado con la matriz de permisos y datos clínicos de demostración relacionados.",
    "Configuración compatible con Railway y documentación operativa de despliegue.",
    "Cobertura funcional de los módulos principales mediante pruebas unitarias y E2E registradas.",
):
    add_bullet(doc, text)

doc.add_heading("4. Objetivos o alcances fallidos", level=1)
add_table(doc, ["Objetivo", "Estado", "Causa y acción correctiva"], [
    ("Instancia Docker en Windows", "En proceso", "Docker Desktop y PostgreSQL todavía deben iniciarse. Seguir 2.2 y sustituir la prueba directa de localhost por la del contenedor."),
    ("Instancia pública Railway", "Pendiente", "Debe ejecutarse después de validar Docker local. Crear proyecto, configurar secretos y registrar la URL."),
    ("Pruebas HTTP en ambas URLs", "Pendiente", "Dependen de que ambas instancias estén activas. Ejecutar la matriz y anexar respuestas/capturas."),
    ("Números de control completos", "Parcial", "Emmanuel y Orlando fueron proporcionados como 2329XXXX; sustituirlos si la versión final exige el número completo."),
], [2550, 1350, 5460])

doc.add_heading("Conclusión general", level=1)
doc.add_paragraph(
    "El proceso de publicación de MediTrackCare quedó explicado desde la instalación de Docker Desktop en Windows hasta el despliegue en Railway. Docker permite ejecutar la API y PostgreSQL sin instalarlos por separado en Windows. Prisma crea y actualiza las tablas, la carga inicial prepara los roles y el administrador, y /api/health ayuda a comprobar si la aplicación está activa. El proyecto compiló y aprobó 40 pruebas unitarias; también existe un resultado anterior de 157 pruebas E2E. Aun así, falta levantar los contenedores, comprobar que los datos se conservan y repetir las pruebas en Railway. Por esta razón el documento separa lo que ya se comprobó de lo que todavía debe ejecutarse."
)

doc.add_page_break()
doc.add_heading("Aprendizaje individual", level=1)
p = doc.add_paragraph()
r = p.add_run("Nombre del integrante: Lisbeth Catalan Castro")
font(r, 11, True, NAVY)
learning = (
    "Durante la colaboración en MediTrackCare comprendí que programar las funciones es solo una parte de un proyecto web. Antes pensaba que desplegar era ejecutar un comando para poner la aplicación en Internet. Ahora entiendo que también hay que preparar la computadora, configurar la base de datos, proteger las contraseñas, hacer pruebas y revisar errores. Preparar una versión en Docker Desktop y otra en Railway me ayudó a entender que el servicio debe funcionar de manera parecida tanto en nuestra computadora con Windows como en la nube.\n\n"
    "Uno de los aprendizajes más importantes fue el uso de Docker como contrato de ejecución. Definir la versión de Node, instalar dependencias, generar el cliente de Prisma, compilar NestJS y declarar el comando de inicio elimina muchas diferencias entre equipos. Docker Compose agregó otra idea importante: la aplicación no trabaja aislada, sino que depende de PostgreSQL y debe esperar a que la base de datos esté lista. También comprendí la utilidad de un volumen persistente, ya que recrear un contenedor no debería borrar la información clínica. La comprobación de salud permite que Docker o Railway detecten si el proceso responde, lo cual resulta más útil que asumir que todo funciona porque el contenedor aparece encendido.\n\n"
    "La colaboración fue importante para repartir el trabajo. Una persona puede revisar los módulos y las rutas, otra puede preparar Docker, otra ejecutar las pruebas y otra reunir las capturas. Aun así, dividir las tareas no significa trabajar por separado. Todos debemos acordar las variables, los puertos, las cuentas de prueba, los resultados esperados y la versión del código. Aprendí que los archivos de despliegue deben guardarse en el repositorio y explicarse para que cualquier integrante pueda repetir los pasos. Una guía clara reduce la dependencia de la persona que hizo la primera configuración.\n\n"
    "También entendí mejor el uso de Prisma. Las migraciones crean los cambios de la base de datos sin tener que hacer cada tabla a mano. La carga inicial crea los roles y la cuenta administrativa. Esta carga se puede ejecutar más de una vez sin repetir la misma información. También aprendí que las contraseñas reales no deben escribirse dentro del repositorio. Esos datos se colocan en variables de entorno y se deben ocultar en capturas, registros y documentos.\n\n"
    "Las pruebas cambiaron mi manera de evaluar el avance. Una compilación exitosa demuestra que el código puede producirse, pero no demuestra que PostgreSQL esté disponible ni que todas las rutas respondan correctamente. Las pruebas unitarias validan componentes específicos y las E2E ejercitan flujos completos; las pruebas contra cada instancia confirman además que red, variables, migraciones y configuración son correctas. Por eso es importante registrar solicitud, código HTTP, respuesta, fecha y evidencia, y no limitarse a escribir que una prueba pasó. También aprendí a distinguir evidencia actual de resultados históricos, porque reutilizar cifras anteriores sin aclararlo podría conducir a conclusiones incorrectas.\n\n"
    "Finalmente, esta actividad me enseñó que colaborar también implica comunicar límites con honestidad. Si no existe acceso a un VPS o el motor de Docker no está disponible, lo correcto es documentar el impedimento, dejar la configuración lista y acordar quién completará la ejecución. Ocultar un pendiente debilita al equipo; hacerlo visible permite resolverlo. Considero que el mayor aprendizaje fue integrar desarrollo y operación como una sola responsabilidad compartida. El proyecto no termina cuando el código funciona en una computadora: termina cuando puede desplegarse de forma repetible, proteger sus datos, responder a pruebas verificables y ser comprendido por los demás integrantes. Esta visión me ayudará a participar de manera más ordenada, segura y responsable en futuros proyectos de software."
)
for block in learning.split("\n\n"):
    doc.add_paragraph(block)
word_count = len(learning.replace("\n", " ").split())
p = doc.add_paragraph()
r = p.add_run(f"Extensión del texto individual: {word_count} palabras.")
font(r, 9.5, True, GRAY)

doc.add_heading("Verificación de criterios de evaluación", level=1)
add_table(doc, ["Criterio", "Estado antes del envío", "Acción final"], [
    ("Aprendizaje individual (mín. 500 palabras)", f"Cumple en estructura: {word_count} palabras.", "Revisar y personalizar con la experiencia real de cada integrante."),
    ("Redacción y orden", "Cumple; contiene secciones, comandos, pruebas y anexos.", "Sustituir los dos números de control incompletos."),
    ("Servidor local y pruebas", "Procedimiento completo; Docker pendiente.", "Ejecutar Docker Desktop en Windows y anexar evidencia de las funciones."),
    ("Servidor VPS y pruebas", "Configuración lista; ejecución pendiente.", "Desplegar en Railway y anexar URL/evidencias."),
    ("Alcances y objetivos fallidos", "Cumple; se documentan logros, pendientes y causas.", "Actualizar estados tras los despliegues."),
    ("Envío vía SAV", "Pendiente del equipo.", "Un integrante debe enviar en la última sesión de la semana 14."),
], [2800, 3000, 3560])

doc.add_heading("Anexo A. Archivos de liberación", level=1)
for text in (
    "Dockerfile: construcción y ejecución de la API en producción.",
    "docker-compose.yml: API, PostgreSQL, persistencia y comprobaciones de salud.",
    ".env.production.example: catálogo seguro de variables requeridas.",
    "railway.toml: construcción, salud y política de reinicio en Railway.",
    "prisma/seed.ts: roles y administrador inicial configurable.",
    "docs/DESPLIEGUE.md: guía operativa y matriz mínima de pruebas.",
):
    add_bullet(doc, text)

doc.add_heading("Anexo B. Evidencias que deben anexarse", level=1)
for text in (
    "Captura de Docker Desktop en ejecución, wsl --version, docker version y docker compose version.",
    "docker compose ps con api y db activos/saludables.",
    "Logs de api mostrando migraciones, semilla y arranque sin exponer secretos.",
    "Respuesta local de /api/health y Swagger en http://localhost:3000/docs.",
    "Resultados de autenticación, 401, 403 y CRUD por módulo.",
    "Resultado nuevo de pruebas E2E con TEST_DATABASE_URL separada.",
    "Panel Railway con servicios backend y PostgreSQL; variables con valores ocultos.",
    "Deploy Logs, dominio público, /api/health y Swagger en Railway.",
):
    add_bullet(doc, text)

doc.add_heading("Anexo C. Fuentes técnicas consultadas", level=1)
for text in (
    "Docker Docs. Install Docker Desktop on Windows. https://docs.docker.com/desktop/setup/install/windows-install/ (consultado el 24/08/2026).",
    "Microsoft Learn. Instalación de WSL. https://learn.microsoft.com/windows/wsl/install (consultado el 24/08/2026).",
    "Railway Docs. Dockerfiles. https://docs.railway.com/builds/dockerfiles (consultado el 24/08/2026).",
    "Railway Docs. PostgreSQL. https://docs.railway.com/databases/postgresql (consultado el 24/08/2026).",
    "Railway Guides. Deploy a Docker Compose App to Production. https://docs.railway.com/guides/docker-compose (consultado el 24/08/2026).",
    "Railway Docs. Healthchecks. https://docs.railway.com/deployments/healthchecks (consultado el 24/08/2026).",
):
    add_bullet(doc, text)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = "Reporte de despliegue de MediTrackCare"
doc.core_properties.subject = "Liberación del servicio web"
doc.core_properties.author = "Equipo MediTrackCare"
doc.core_properties.keywords = "NestJS, Docker, Railway, Prisma, PostgreSQL"
doc.save(OUT)
print(OUT)
print(f"Palabras del aprendizaje individual: {word_count}")
